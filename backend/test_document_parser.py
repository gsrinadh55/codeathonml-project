"""Unit tests for PlagiSense document parser."""

import os
from pathlib import Path
import tempfile
import unittest

import docx
from backend.document_parser import extract_text_from_file, SUPPORTED_EXTENSIONS


def _create_minimal_pdf(pages_text: list[str]) -> bytes:
    """Create a minimal valid raw PDF 1.4 binary containing the specified pages of text."""
    # Build PDF objects
    objects = []
    
    # 1: Catalog
    # 2: Pages tree
    # 3..3+N-1: Page objects
    # 3+N..3+2N-1: Content stream objects
    # Font object
    num_pages = len(pages_text)
    page_obj_ids = [3 + i for i in range(num_pages)]
    content_obj_ids = [3 + num_pages + i for i in range(num_pages)]
    font_obj_id = 3 + 2 * num_pages

    kids_str = " ".join(f"{pid} 0 R" for pid in page_obj_ids)
    
    # Obj 1: Catalog
    obj_1 = "1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
    # Obj 2: Pages
    obj_2 = f"2 0 obj\n<< /Type /Pages /Kids [{kids_str}] /Count {num_pages} >>\nendobj\n"
    
    page_objs = []
    content_objs = []
    
    for i, text in enumerate(pages_text):
        p_id = page_obj_ids[i]
        c_id = content_obj_ids[i]
        # Escape parenthesis in text
        safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_content = f"BT\n/F1 12 Tf\n50 700 Td\n({safe_text}) Tj\nET"
        stream_bytes = stream_content.encode("latin-1")
        
        page_objs.append(
            f"{p_id} 0 obj\n"
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {c_id} 0 R /Resources << /Font << /F1 {font_obj_id} 0 R >> >> >>\n"
            f"endobj\n"
        )
        content_objs.append(
            f"{c_id} 0 obj\n"
            f"<< /Length {len(stream_bytes)} >>\n"
            f"stream\n{stream_content}\nendstream\n"
            f"endobj\n"
        )

    font_obj = f"{font_obj_id} 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    
    header = b"%PDF-1.4\n"
    body_parts = [obj_1, obj_2] + page_objs + content_objs + [font_obj]
    
    pdf_bytes = bytearray(header)
    offsets = [0]
    
    for part in body_parts:
        offsets.append(len(pdf_bytes))
        pdf_bytes.extend(part.encode("latin-1"))
        
    xref_offset = len(pdf_bytes)
    total_objects = font_obj_id + 1
    
    xref = f"xref\n0 {total_objects}\n0000000000 65535 f \n"
    for off in offsets[1:]:
        xref += f"{off:010d} 00000 n \n"
        
    trailer = (
        f"trailer\n"
        f"<< /Size {total_objects} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    )
    pdf_bytes.extend(trailer.encode("latin-1"))
    return bytes(pdf_bytes)


class TestDocumentParser(unittest.TestCase):
    """Test suite for document_parser extraction and error handling."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    # --- TXT TESTS ---

    def test_extract_txt_utf8(self):
        """Test extraction from standard UTF-8 text file with whitespace normalization."""
        txt_file = self.dir_path / "sample_utf8.txt"
        content = "  Line 1: Plagiarism detection engine.   \n\n\n\n   Line 2: Academic integrity.  \n"
        txt_file.write_text(content, encoding="utf-8")

        result = extract_text_from_file(str(txt_file))
        expected = "Line 1: Plagiarism detection engine.\n\nLine 2: Academic integrity."
        self.assertEqual(result, expected)

    def test_extract_txt_encodings(self):
        """Test fallback decoding for cp1252 / latin-1 encoded text."""
        txt_file = self.dir_path / "sample_latin1.txt"
        latin_text = "Café résumé with copyright © 2026."
        txt_file.write_bytes(latin_text.encode("latin-1"))

        result = extract_text_from_file(str(txt_file))
        self.assertIn("Café", result)
        self.assertIn("résumé", result)
        self.assertIn("© 2026", result)

    def test_extract_txt_empty_and_whitespace(self):
        """Test extraction from 0-byte and whitespace-only TXT files."""
        empty_file = self.dir_path / "empty.txt"
        empty_file.write_text("", encoding="utf-8")
        self.assertEqual(extract_text_from_file(str(empty_file)), "")

        whitespace_file = self.dir_path / "whitespace.txt"
        whitespace_file.write_text("   \n\n   \t  \n   ", encoding="utf-8")
        self.assertEqual(extract_text_from_file(str(whitespace_file)), "")

    # --- PDF TESTS ---

    def test_extract_pdf_multipage(self):
        """Test extraction from a multi-page PDF document."""
        pdf_file = self.dir_path / "multipage.pdf"
        pages = ["First page text about plagiarism.", "Second page text about originality."]
        pdf_bytes = _create_minimal_pdf(pages)
        pdf_file.write_bytes(pdf_bytes)

        result = extract_text_from_file(str(pdf_file))
        self.assertIn("First page text about plagiarism.", result)
        self.assertIn("Second page text about originality.", result)

    def test_extract_pdf_empty_and_zero_byte(self):
        """Test 0-byte PDF handling."""
        empty_pdf = self.dir_path / "empty.pdf"
        empty_pdf.write_bytes(b"")
        self.assertEqual(extract_text_from_file(str(empty_pdf)), "")

    def test_extract_pdf_corrupted(self):
        """Test that corrupted PDF raises ValueError."""
        corrupt_pdf = self.dir_path / "corrupt.pdf"
        corrupt_pdf.write_bytes(b"This is not a valid PDF content header %PDF broken junk")
        with self.assertRaises(ValueError) as ctx:
            extract_text_from_file(str(corrupt_pdf))
        self.assertIn("corrupted or invalid PDF", str(ctx.exception))

    # --- DOCX TESTS ---

    def test_extract_docx_paragraphs(self):
        """Test extraction from DOCX with multiple paragraphs and blank lines."""
        docx_file = self.dir_path / "sample.docx"
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1: Introduction to Natural Language Processing.")
        doc.add_paragraph("")  # empty paragraph
        doc.add_paragraph("   ")  # whitespace paragraph
        doc.add_paragraph("Paragraph 2: Evaluating similarity across documents.")
        doc.save(str(docx_file))

        result = extract_text_from_file(str(docx_file))
        expected = "Paragraph 1: Introduction to Natural Language Processing.\n\nParagraph 2: Evaluating similarity across documents."
        self.assertEqual(result, expected)

    def test_extract_docx_empty_and_zero_byte(self):
        """Test extraction from empty DOCX and 0-byte DOCX."""
        empty_docx = self.dir_path / "empty_doc.docx"
        doc = docx.Document()
        doc.save(str(empty_docx))
        self.assertEqual(extract_text_from_file(str(empty_docx)), "")

        zero_byte = self.dir_path / "zero.docx"
        zero_byte.write_bytes(b"")
        self.assertEqual(extract_text_from_file(str(zero_byte)), "")

    def test_extract_docx_corrupted(self):
        """Test that corrupted DOCX file raises ValueError."""
        corrupt_docx = self.dir_path / "corrupt.docx"
        corrupt_docx.write_bytes(b"Not a zip or docx file byte stream")
        with self.assertRaises(ValueError) as ctx:
            extract_text_from_file(str(corrupt_docx))
        self.assertIn("corrupted or invalid DOCX", str(ctx.exception))

    # --- ERROR & EDGE CASE TESTS ---

    def test_unsupported_extensions(self):
        """Test that unsupported file extensions are rejected with clear ValueError."""
        unsupported_files = ["doc.png", "data.csv", "table.xlsx", "notes.md", "archive.zip"]
        for filename in unsupported_files:
            file_path = self.dir_path / filename
            file_path.write_text("dummy content")
            with self.assertRaises(ValueError) as ctx:
                extract_text_from_file(str(file_path))
            self.assertIn("Unsupported file format", str(ctx.exception))
            self.assertIn(".txt", str(ctx.exception))
            self.assertIn(".pdf", str(ctx.exception))
            self.assertIn(".docx", str(ctx.exception))

    def test_nonexistent_file(self):
        """Test that non-existent file path raises FileNotFoundError."""
        with self.assertRaises(FileNotFoundError):
            extract_text_from_file(str(self.dir_path / "non_existent_document.txt"))


if __name__ == "__main__":
    unittest.main()
